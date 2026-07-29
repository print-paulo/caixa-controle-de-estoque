import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.conectar_banco import conectar_banco
from models.venda import Venda
from models.compra import Compra
from services.estoque import listar_estoque_completo, listar_movimentos
from services.relatorios import (
    relatorio_produtos, relatorio_estoque, relatorio_vendas,
    relatorio_compras, relatorio_lucro,
)

# ---------- estilo (visual consistente em todas as abas) ----------

FONTE = "Calibri"
COR_CABECALHO_FUNDO = "1F4E78"   # azul escuro
COR_CABECALHO_TEXTO = "FFFFFF"   # branco
COR_TITULO_FUNDO = "D9E1F2"      # azul claro (título de seção no Resumo)
COR_ALERTA_FUNDO = "FCE4E4"      # vermelho claro (estoque abaixo do mínimo)

FORMATO_MOEDA = 'R$ #,##0.00;[RED]-R$ #,##0.00'
FORMATO_PORCENTAGEM = '0.0%'

_BORDA_FINA = Side(style="thin", color="B7B7B7")
BORDA_CELULA = Border(left=_BORDA_FINA, right=_BORDA_FINA, top=_BORDA_FINA, bottom=_BORDA_FINA)


def _estilo_cabecalho(ws, linha, ultima_coluna):
    """Aplica o estilo de cabeçalho (fundo azul, texto branco em negrito) numa linha inteira."""
    for col in range(1, ultima_coluna + 1):
        celula = ws.cell(row=linha, column=col)
        celula.font = Font(name=FONTE, bold=True, color=COR_CABECALHO_TEXTO)
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO_FUNDO)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        celula.border = BORDA_CELULA


def _escrever_tabela(ws, cabecalhos, linhas, linha_inicial=1, formatos_coluna=None):
    """
    Escreve uma tabela (cabeçalho + linhas de dados) a partir da linha_inicial,
    aplica o estilo de cabeçalho, borda em toda célula de dado, congela o
    cabeçalho e ajusta a largura de cada coluna pelo maior conteúdo.

    - formatos_coluna: dict opcional {índice_da_coluna (1-based): number_format}
    """
    formatos_coluna = formatos_coluna or {}

    for col, texto in enumerate(cabecalhos, start=1):
        ws.cell(row=linha_inicial, column=col, value=texto)
    _estilo_cabecalho(ws, linha_inicial, len(cabecalhos))

    for i, linha_dados in enumerate(linhas, start=linha_inicial + 1):
        for col, valor in enumerate(linha_dados, start=1):
            celula = ws.cell(row=i, column=col, value=valor)
            celula.font = Font(name=FONTE)
            celula.border = BORDA_CELULA
            if col in formatos_coluna:
                celula.number_format = formatos_coluna[col]

    ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=1)
    _autoajustar_largura(ws, len(cabecalhos))

    return linha_inicial + len(linhas)  # última linha escrita


def _autoajustar_largura(ws, num_colunas, largura_minima=10, largura_maxima=45):
    for col in range(1, num_colunas + 1):
        letra = get_column_letter(col)
        maior = largura_minima
        for celula in ws[letra]:
            if celula.value is not None:
                maior = max(maior, len(str(celula.value)) + 2)
        ws.column_dimensions[letra].width = min(maior, largura_maxima)


# ---------- período (mesmo formato dos relatórios existentes) ----------

def _dentro_do_periodo(data_hora, data_inicio, data_fim):
    """data_hora no formato 'YYYY-MM-DD HH:MM:SS'; data_inicio/data_fim em 'YYYY-MM-DD'."""
    data = data_hora[:10]
    if data_inicio and data < data_inicio:
        return False
    if data_fim and data > data_fim:
        return False
    return True


def _listar_vendas_periodo(data_inicio, data_fim):
    conn = conectar_banco()
    try:
        rows = conn.execute("""
            SELECT * FROM venda
            WHERE status = 'FINALIZADA'
            ORDER BY data_hora DESC, id_venda DESC
        """).fetchall()
        vendas = Venda.from_rows(rows)
        return [v for v in vendas if _dentro_do_periodo(v.data_hora, data_inicio, data_fim)]
    finally:
        conn.close()


def _listar_compras_periodo(data_inicio, data_fim):
    conn = conectar_banco()
    try:
        rows = conn.execute("""
            SELECT * FROM compra
            WHERE status = 'FINALIZADA'
            ORDER BY data_hora DESC, id_compra DESC
        """).fetchall()
        compras = Compra.from_rows(rows)
        return [c for c in compras if _dentro_do_periodo(c.data_hora, data_inicio, data_fim)]
    finally:
        conn.close()


def _listar_movimentos_periodo(data_inicio, data_fim):
    todos = listar_movimentos(limite=1_000_000)
    return [m for m in todos if _dentro_do_periodo(m.data_hora, data_inicio, data_fim)]


# ---------- abas ----------

def _montar_aba_resumo(wb, data_inicio, data_fim):
    ws = wb.active
    ws.title = "Resumo"

    produtos = relatorio_produtos()
    estoque = relatorio_estoque()
    vendas = relatorio_vendas(data_inicio, data_fim)
    compras = relatorio_compras(data_inicio, data_fim)
    lucro = relatorio_lucro(data_inicio, data_fim)

    periodo_texto = "Todo o histórico"
    if data_inicio or data_fim:
        periodo_texto = f"{data_inicio or '(início)'} até {data_fim or '(hoje)'}"

    linha = 1
    ws.cell(row=linha, column=1, value="Relatório do Caixa & Controle de Estoque").font = Font(
        name=FONTE, bold=True, size=16, color=COR_CABECALHO_FUNDO
    )
    linha += 1
    ws.cell(row=linha, column=1, value=f"Período: {periodo_texto}").font = Font(name=FONTE, italic=True)
    linha += 2

    def secao(titulo):
        nonlocal linha
        ws.cell(row=linha, column=1, value=titulo)
        ws.cell(row=linha, column=1).font = Font(name=FONTE, bold=True, size=12)
        ws.cell(row=linha, column=1).fill = PatternFill("solid", fgColor=COR_TITULO_FUNDO)
        ws.cell(row=linha, column=2).fill = PatternFill("solid", fgColor=COR_TITULO_FUNDO)
        linha += 1

    def linha_kv(rotulo, valor, formato=None):
        nonlocal linha
        ws.cell(row=linha, column=1, value=rotulo).font = Font(name=FONTE)
        celula_valor = ws.cell(row=linha, column=2, value=valor)
        celula_valor.font = Font(name=FONTE, bold=True)
        if formato:
            celula_valor.number_format = formato
        linha += 1

    secao("Produtos")
    linha_kv("Produtos ativos", produtos["total_ativos"])
    linha_kv("Produtos desativados", produtos["total_inativos"])
    linha_kv("Produtos sem preço cadastrado", len(produtos["sem_preco"]))
    linha += 1

    secao("Estoque")
    linha_kv("Valor total em estoque", estoque["valor_total_estoque"], FORMATO_MOEDA)
    linha_kv("Unidades totais em estoque", estoque["quantidade_total_unidades"])
    linha_kv("Produtos no mínimo ou abaixo", estoque["produtos_abaixo_minimo"])
    linha += 1

    secao("Vendas no período")
    linha_kv("Quantidade de vendas finalizadas", vendas["quantidade_vendas"])
    linha_kv("Total vendido", vendas["total_vendido"], FORMATO_MOEDA)
    linha_kv("Ticket médio", vendas["ticket_medio"], FORMATO_MOEDA)
    linha += 1

    secao("Compras no período")
    linha_kv("Quantidade de compras finalizadas", compras["quantidade_compras"])
    linha_kv("Total investido", compras["total_investido"], FORMATO_MOEDA)
    linha_kv("Lucro esperado (se tudo for vendido pelo preço calculado)", compras["lucro_esperado_total"], FORMATO_MOEDA)
    linha += 1

    secao("Lucro no período")
    linha_kv("Lucro bruto (regime de caixa: vendido - investido)", lucro["lucro_bruto"], FORMATO_MOEDA)
    linha_kv("Lucro real (margem das vendas já realizadas)", lucro["lucro_real"], FORMATO_MOEDA)
    if lucro["unidades_sem_custo_registrado"]:
        linha_kv("Unidades vendidas sem custo registrado (fora do lucro real)", lucro["unidades_sem_custo_registrado"])

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20


def _montar_aba_produtos(wb):
    ws = wb.create_sheet("Produtos")
    conn = conectar_banco()
    try:
        rows = conn.execute("""
            SELECT p.id_produto, p.nome_produto, p.codigo_barras, c.nome_categoria,
                   p.medida_embalagem, p.unidade, p.valor_unitario, p.custo_unitario
            FROM produto p
            LEFT JOIN categoria c ON c.id_categoria = p.id_categoria
            WHERE p.ativo = 1
            ORDER BY p.nome_produto
        """).fetchall()
    finally:
        conn.close()

    cabecalhos = ["ID", "Nome", "Código de Barras", "Categoria", "Medida", "Unidade", "Preço de Venda", "Custo Unitário"]
    linhas = [
        (r["id_produto"], r["nome_produto"], r["codigo_barras"] or "—", r["nome_categoria"] or "—",
         r["medida_embalagem"] or "—", r["unidade"] or "—", r["valor_unitario"], r["custo_unitario"])
        for r in rows
    ]
    _escrever_tabela(ws, cabecalhos, linhas, formatos_coluna={7: FORMATO_MOEDA, 8: FORMATO_MOEDA})


def _montar_aba_estoque(wb):
    ws = wb.create_sheet("Estoque")
    estoque = listar_estoque_completo()
    cabecalhos = ["ID Produto", "Produto", "Depósito", "Exposição", "Capacidade Exposição", "Mínimo", "Última Atualização"]
    linhas = [
        (e.id_produto, e.nome_produto, e.estoque_deposito, e.estoque_exposicao,
         e.capacidade_exposicao, e.estoque_minimo, e.ultima_atualizacao)
        for e in estoque
    ]
    ultima_linha = _escrever_tabela(ws, cabecalhos, linhas)

    # destaca em vermelho claro as linhas de produto no mínimo ou abaixo dele
    for i, e in enumerate(estoque, start=2):
        if e.estoque_minimo is not None and e.estoque_deposito <= e.estoque_minimo:
            for col in range(1, len(cabecalhos) + 1):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=COR_ALERTA_FUNDO)


def _montar_aba_vendas(wb, data_inicio, data_fim):
    ws = wb.create_sheet("Vendas")
    vendas = _listar_vendas_periodo(data_inicio, data_fim)
    cabecalhos = ["ID", "Data/Hora", "Forma de Pagamento", "Valor Total"]
    linhas = [(v.id_venda, v.data_hora, v.forma_pagamento, v.valor_total) for v in vendas]
    _escrever_tabela(ws, cabecalhos, linhas, formatos_coluna={4: FORMATO_MOEDA})


def _montar_aba_compras(wb, data_inicio, data_fim):
    ws = wb.create_sheet("Compras")
    compras = _listar_compras_periodo(data_inicio, data_fim)
    cabecalhos = ["ID", "Data/Hora", "Fornecedor", "Status"]
    linhas = [(c.id_compra, c.data_hora, c.fornecedor or "—", c.status) for c in compras]
    _escrever_tabela(ws, cabecalhos, linhas)


def _montar_aba_movimentacao(wb, data_inicio, data_fim):
    ws = wb.create_sheet("Movimentação")
    movimentos = _listar_movimentos_periodo(data_inicio, data_fim)
    cabecalhos = ["Data/Hora", "Produto", "Tipo", "Campo", "Quantidade", "Origem (venda/compra)"]
    linhas = [
        (m.data_hora, m.nome_produto, m.tipo, m.campo, m.quantidade, m.origem_id or "—")
        for m in movimentos
    ]
    _escrever_tabela(ws, cabecalhos, linhas)


# ---------- função principal ----------

def exportar_relatorio_excel(caminho_arquivo, data_inicio=None, data_fim=None):
    """
    Gera um arquivo .xlsx com o snapshot completo do sistema: um resumo
    executivo e uma aba por entidade (Produtos, Estoque, Vendas, Compras,
    Movimentação), formatado (cabeçalho colorido, moeda em R$, largura de
    coluna ajustada, estoque abaixo do mínimo destacado).

    - data_inicio/data_fim (formato 'YYYY-MM-DD', ambos opcionais): filtram
      vendas, compras e movimentação de estoque pelo período. Produtos e
      Estoque são sempre o snapshot ATUAL (não faz sentido "estoque no
      passado" sem rastrear histórico ponto-a-ponto).

    Levanta ValueError se data_inicio for depois de data_fim.
    Retorna o caminho do arquivo gerado.
    """
    if data_inicio and data_fim and data_inicio > data_fim:
        raise ValueError(f"Data de início ({data_inicio}) não pode ser depois da data de fim ({data_fim}).")

    wb = Workbook()
    _montar_aba_resumo(wb, data_inicio, data_fim)
    _montar_aba_produtos(wb)
    _montar_aba_estoque(wb)
    _montar_aba_vendas(wb, data_inicio, data_fim)
    _montar_aba_compras(wb, data_inicio, data_fim)
    _montar_aba_movimentacao(wb, data_inicio, data_fim)

    caminho_arquivo = str(caminho_arquivo)
    if not caminho_arquivo.lower().endswith(".xlsx"):
        caminho_arquivo += ".xlsx"

    wb.save(caminho_arquivo)
    return caminho_arquivo