import openpyxl
import pytest

from services.registrar_compra import iniciar_compra, adicionar_item_compra, finalizar_compra
from services.registrar_venda import iniciar_venda, adicionar_item_venda, finalizar_venda
from services.exportar_excel import exportar_relatorio_excel, FORMATO_MOEDA


def _comprar(codigo, quantidade, custo=30.0, margem=0.3, fornecedor="Fornecedor Teste"):
    id_compra = iniciar_compra(fornecedor)
    adicionar_item_compra(id_compra, codigo, quantidade, custo, margem)
    finalizar_compra(id_compra)
    return id_compra


def _vender(codigo, quantidade, forma_pagamento="dinheiro"):
    id_venda = iniciar_venda()
    adicionar_item_venda(id_venda, codigo, quantidade)
    finalizar_venda(id_venda, forma_pagamento)
    return id_venda


class TestExportarRelatorioExcel:
    def test_gera_arquivo_com_todas_as_abas_esperadas(self, produto_padrao, tmp_path):
        caminho = exportar_relatorio_excel(tmp_path / "relatorio.xlsx")
        wb = openpyxl.load_workbook(caminho)
        assert wb.sheetnames == ["Resumo", "Produtos", "Estoque", "Vendas", "Compras", "Movimentação"]

    def test_adiciona_extensao_xlsx_se_faltar(self, produto_padrao, tmp_path):
        caminho = exportar_relatorio_excel(tmp_path / "sem_extensao")
        assert caminho.endswith(".xlsx")
        assert openpyxl.load_workbook(caminho)  # não levanta

    def test_rejeita_periodo_com_inicio_depois_do_fim(self, tmp_path):
        with pytest.raises(ValueError, match="não pode ser depois"):
            exportar_relatorio_excel(tmp_path / "x.xlsx", data_inicio="2026-12-31", data_fim="2026-01-01")

    def test_aba_produtos_lista_produto_com_categoria_e_precos(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 10, custo=25.0, margem=0.5)  # define preço/custo
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Produtos"]

        cabecalho = [c.value for c in ws[1]]
        assert cabecalho == [
            "ID", "Nome", "Código de Barras", "Categoria", "Medida",
            "Unidade", "Preço de Venda", "Custo Unitário",
        ]
        linha = [c.value for c in ws[2]]
        assert linha[1] == "Vinho Tinto Reserva"
        assert linha[3] == "Vinhos"
        assert linha[6] == pytest.approx(37.5)
        assert linha[7] == pytest.approx(25.0)

    def test_aba_produtos_nao_lista_produto_inativo(self, produto_padrao, tmp_path):
        from services.excluir_produto import excluir_produto
        excluir_produto(produto_padrao)

        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Produtos"]
        assert ws.max_row == 1  # só o cabeçalho, produto inativo não aparece

    def test_aba_estoque_reflete_estoque_atual(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20)
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Estoque"]
        linha = [c.value for c in ws[2]]
        assert linha[1] == "Vinho Tinto Reserva"
        assert linha[2] == 20  # depósito (reposição só acontece na venda, não na compra)
        assert linha[3] == 0   # exposição (ainda 0, nenhuma venda ocorreu)

    def test_aba_estoque_destaca_produto_abaixo_do_minimo(self, produto_padrao, tmp_path):
        # produto_padrao tem estoque_minimo=3 e nasce com depósito=0 -> já está abaixo
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Estoque"]
        celula = ws.cell(row=2, column=1)
        assert celula.fill.fgColor.rgb == "00FCE4E4"  # cor de alerta

    def test_aba_estoque_nao_destaca_produto_acima_do_minimo(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 50)  # bem acima do mínimo (3)
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Estoque"]
        celula = ws.cell(row=2, column=1)
        assert celula.fill.fgColor.rgb in ("00000000", None)  # sem preenchimento de alerta

    def test_aba_vendas_lista_apenas_vendas_finalizadas_com_total_e_formato_moeda(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20)
        _vender("7891234567895", 3, "credito")
        iniciar_venda()  # venda em aberto, não deve aparecer

        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Vendas"]
        assert ws.max_row == 2  # cabeçalho + 1 venda finalizada
        linha = [c.value for c in ws[2]]
        assert linha[2] == "credito"
        assert linha[3] > 117.0  # com taxa de cartão de crédito aplicada
        assert ws.cell(row=2, column=4).number_format == FORMATO_MOEDA

    def test_aba_compras_lista_apenas_compras_finalizadas(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20, fornecedor="Fornecedor Único")
        iniciar_compra("Fornecedor Nunca Finalizado")

        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Compras"]
        assert ws.max_row == 2
        linha = [c.value for c in ws[2]]
        assert linha[2] == "Fornecedor Único"
        assert linha[3] == "FINALIZADA"

    def test_aba_movimentacao_lista_historico_do_produto(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20)
        _vender("7891234567895", 3)

        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Movimentação"]
        tipos = {row[2].value for row in ws.iter_rows(min_row=2)}
        assert tipos == {"COMPRA", "REPOSICAO", "VENDA"}

    def test_periodo_filtra_vendas_compras_e_movimentacao_mas_nao_produtos_e_estoque(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20)
        _vender("7891234567895", 3)

        # período no passado, antes de qualquer coisa existir
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx", data_inicio="2000-01-01", data_fim="2000-01-02")
        wb = openpyxl.load_workbook(caminho)

        assert wb["Vendas"].max_row == 1       # só cabeçalho, nada no período
        assert wb["Compras"].max_row == 1
        assert wb["Movimentação"].max_row == 1
        assert wb["Produtos"].max_row == 2     # produto continua aparecendo (snapshot atual)
        assert wb["Estoque"].max_row == 2      # estoque também

    def test_resumo_mostra_periodo_e_metricas_principais(self, produto_padrao, tmp_path):
        _comprar("7891234567895", 20, custo=30.0, margem=0.3)
        _vender("7891234567895", 3)

        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx", data_inicio="2020-01-01", data_fim="2030-01-01")
        ws = openpyxl.load_workbook(caminho)["Resumo"]

        valores_coluna_a = [c.value for c in ws["A"] if c.value]
        assert any("2020-01-01" in v and "2030-01-01" in v for v in valores_coluna_a)
        assert "Produtos" in valores_coluna_a
        assert "Vendas no período" in valores_coluna_a
        assert "Lucro no período" in valores_coluna_a

    def test_resumo_sem_periodo_mostra_todo_o_historico(self, produto_padrao, tmp_path):
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        ws = openpyxl.load_workbook(caminho)["Resumo"]
        valores_coluna_a = [c.value for c in ws["A"] if c.value]
        assert any("Todo o histórico" in v for v in valores_coluna_a)

    def test_cabecalho_tem_estilo_de_destaque_em_todas_as_abas_tabulares(self, produto_padrao, tmp_path):
        caminho = exportar_relatorio_excel(tmp_path / "r.xlsx")
        wb = openpyxl.load_workbook(caminho)
        for nome in ["Produtos", "Estoque", "Vendas", "Compras", "Movimentação"]:
            celula = wb[nome].cell(row=1, column=1)
            assert celula.font.bold is True
            assert celula.fill.fgColor.rgb == "001F4E78"